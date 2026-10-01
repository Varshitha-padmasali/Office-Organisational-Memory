"""
Unit tests for extraction_service.extract_text().

.txt and .docx tests do a real round-trip (write a real file, extract it
back) using only the stdlib and python-docx, which is already a backend
dependency. The .pdf case is tested via monkeypatching pypdf.PdfReader,
since generating a real text-bearing PDF would require an extra dependency
(e.g. reportlab) not otherwise needed by this project — we're testing our
own dispatch/joining logic here, not pypdf's PDF parsing.
"""

from pathlib import Path

import pytest

from app.services.extraction_service import ExtractionError, extract_text


def test_extract_text_from_txt_file(tmp_path: Path):
    file_path = tmp_path / "notes.txt"
    file_path.write_text("Hello from a plain text file.", encoding="utf-8")

    result = extract_text(str(file_path), "notes.txt")

    assert result == "Hello from a plain text file."


def test_extract_text_from_docx_file(tmp_path: Path):
    from docx import Document as DocxDocument

    file_path = tmp_path / "notes.docx"
    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    doc.save(str(file_path))

    result = extract_text(str(file_path), "notes.docx")

    assert "First paragraph." in result
    assert "Second paragraph." in result


def test_extract_text_from_pdf_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    file_path = tmp_path / "notes.pdf"
    file_path.write_bytes(b"%PDF-1.4 fake content for the test")

    class FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        def __init__(self, path: str) -> None:
            self.pages = [FakePage("Page one text."), FakePage("Page two text.")]

    monkeypatch.setattr("app.services.extraction_service.PdfReader", FakeReader)

    result = extract_text(str(file_path), "notes.pdf")

    assert "Page one text." in result
    assert "Page two text." in result


def test_extract_text_rejects_unsupported_extension(tmp_path: Path):
    file_path = tmp_path / "notes.exe"
    file_path.write_bytes(b"not a real document")

    with pytest.raises(ExtractionError):
        extract_text(str(file_path), "notes.exe")


def test_extract_text_rejects_empty_txt_file(tmp_path: Path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("   \n\t  ", encoding="utf-8")

    with pytest.raises(ExtractionError):
        extract_text(str(file_path), "empty.txt")
